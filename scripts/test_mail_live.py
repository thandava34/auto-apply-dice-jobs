"""Opt-in, one-message live acceptance harness. Never runs under pytest.

Uses the production adapter with a test-only config, explicit test-only CC/BCC, and a
generated non-personal attachment. Default action is a read-only readiness probe.
Run one provider at a time; authentication must be completed by the operator.
"""
import argparse
import base64
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
import hashlib
import json
from pathlib import Path
import time
import uuid

from core.outreach.email_engine import EmailEngine
from core.outreach.mail_contract import address_list, verify_mime


URLS = {
    "outlook_web": "https://outlook.office.com/mail/",
    "gmail_web": "https://mail.google.com/mail/u/0/#inbox",
    "zoho_web": "https://mail.zoho.com/zm/#mail/folder/inbox",
}
COMPOSE = {
    "outlook_web": "//button[@aria-label='New mail' or @title='New mail' or @aria-label='New email' or @title='New email' or normalize-space(.)='New email' or @data-unique-id='Ribbon-NewMail']",
    "gmail_web": "//*[@gh='cm' or @data-tooltip='Compose'] | //div[@role='button' and normalize-space(.)='Compose']",
    "zoho_web": "//*[@id='new_mail'] | //*[@title='New Mail'] | //button[normalize-space(.)='New Mail']",
}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--provider", required=True, choices=[*URLS, "gmail_api", "outlook"])
    parser.add_argument("--recipient", required=True, help="One operator-owned test address")
    parser.add_argument("--cc", default="", help="Explicitly approved test CC addresses only")
    parser.add_argument("--bcc", default="", help="Explicitly approved test BCC addresses only")
    parser.add_argument("--mode", choices=["probe", "draft", "send"], default="probe")
    parser.add_argument("--login-wait", type=int, default=45)
    parser.add_argument("--hold", action="store_true", help="Keep owned test browser open for inspection until Enter")
    parser.add_argument("--run-id", default="", help="Unique ID; an existing run directory is never reused")
    args = parser.parse_args()
    recipients = address_list(args.recipient)
    cc = ",".join(address_list(args.cc))
    bcc = ",".join(address_list(args.bcc))
    if len(recipients) != 1:
        parser.error("Exactly one operator-owned recipient is required")
    run_id = args.run_id or (datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid.uuid4().hex[:6])
    if not run_id.replace("-", "").isalnum():
        parser.error("run-id must contain only letters, digits and hyphens")
    output = Path("data/mail_tests") / run_id
    output.mkdir(parents=True, exist_ok=False)
    report = {"run_id": run_id, "provider": args.provider, "mode": args.mode,
              "recipient": recipients[0], "cc": cc, "bcc": bcc,
              "result": "not_started", "delivery_verified": False}

    def checkpoint():
        (output / "result.json").write_text(json.dumps(report, indent=2), encoding="utf-8")

    config_path = Path("config/outreach_settings.json")
    config = json.loads(config_path.read_text(encoding="utf-8")) if config_path.exists() else {}
    config.update(email_provider=args.provider, send_mode="send" if args.mode == "send" else "draft",
                  cc_email="", bcc_email="", use_ai_email=False, require_resume_attachment=True,
                  excluded_vendor_domains="", excluded_email_addresses="", attachment_timeout_seconds=45)
    engine = EmailEngine(config)
    engine.log = lambda value: print(str(value), flush=True)
    driver = None
    checkpoint()
    try:
        if args.provider in URLS:
            driver = engine._get_web_driver(args.provider, URLS[args.provider], log_fn=engine.log)
            deadline = time.monotonic() + max(1, args.login_wait)
            print("Browser opened. If sign-in is required, complete it yourself; no credentials are entered by this script.", flush=True)
            while time.monotonic() < deadline:
                from urllib.parse import urlsplit
                page = urlsplit(driver.current_url)
                if page.hostname == "accounts.google.com" and "/signin/rejected" in page.path:
                    report["result"] = "blocked_google_browser_security_use_gmail_api"
                    driver.save_screenshot(str(output / "browser.png"))
                    return 2
                if any(el.is_displayed() for el in driver.find_elements("xpath", COMPOSE[args.provider])):
                    break
                time.sleep(.5)
            else:
                report["result"] = "blocked_login_or_compose_unavailable"
                report["compose_candidates"] = driver.execute_script("""
                    return [...document.querySelectorAll('button,[role="button"],a')]
                      .filter(e=>e.getClientRects().length && /new (mail|email)|compose/i.test(
                        (e.innerText||'')+' '+(e.getAttribute('aria-label')||'')+' '+(e.title||'')))
                      .map(e=>({tag:e.tagName,role:e.getAttribute('role'),label:e.getAttribute('aria-label'),
                        title:e.title,text:e.innerText,id:e.id}));
                """)
                driver.save_screenshot(str(output / "browser.png"))
                print(json.dumps(report), flush=True)
                return 2
        elif args.provider == "gmail_api":
            service = engine._get_gmail_service()
            service.users().drafts().list(userId="me", maxResults=1).execute()
        else:
            import winreg
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Outlook.Application\CLSID"):
                pass
        if args.mode == "probe":
            report["result"] = "ready_for_controlled_test"
            return 0

        from docx import Document
        attachment = output / "Nvoids test résumé.docx"
        document = Document()
        document.add_heading("Nvoids mail acceptance test", 0)
        document.add_paragraph("Synthetic attachment only. This is not a candidate resume or a job application.")
        document.add_paragraph("Formatting check: café, résumé, Python & SQL <ETL>.")
        document.save(str(attachment))
        payload = attachment.read_bytes()
        report["attachment_sha256"] = hashlib.sha256(payload).hexdigest()
        subject = f"[Nvoids TEST {run_id}] {args.provider} {args.mode}"
        body = ("Hello,\n\nThis is an authorized Nvoids mail test, not a job application.\n"
                "Role: Synthetic Data Engineer\nCompany: Test Company\nLocation: Test Location\n\n"
                "Formatting: café, résumé, Python & SQL <ETL>.\n"
                "The attached document contains synthetic test content only.\n\nThanks,\nNvoids test runner")
        job = {"Recruiter Email": recipients[0], "CC": cc, "BCC": bcc}
        report.update(subject=subject, result="action_started_do_not_retry_blindly")
        checkpoint()  # Durable intent before any external draft/send action.
        report["result"] = engine.send_email(job, str(attachment.resolve()), body, subject)
        report["receipt"] = engine.last_provider_receipt
        checkpoint()
        if args.provider == "gmail_api" and report["result"] == "Draft":
            receipt = engine.last_provider_receipt
            saved = engine._get_gmail_service().users().drafts().get(
                userId="me", id=receipt["draft_id"], format="raw").execute()
            raw = saved["message"]["raw"]
            message = BytesParser(policy=policy.default).parsebytes(base64.urlsafe_b64decode(raw + "=" * (-len(raw) % 4)))
            verify_mime(message, job, subject, body, attachment.name, payload)
            report["saved_mime_verified"] = True
        if driver:
            driver.save_screenshot(str(output / "browser.png"))
            # Inspect only controls relevant to the owned test compose, not inbox messages.
            controls = driver.execute_script("""
                return [...document.querySelectorAll('input,textarea,[contenteditable="true"]')]
                  .filter(e=>e.getClientRects().length).map(e=>({tag:e.tagName,name:e.name||'',
                   label:e.getAttribute('aria-label')||'',id:e.id||'',value:e.value||e.innerText||''}));
            """)
            (output / "compose-controls.json").write_text(json.dumps(controls, indent=2), encoding="utf-8")
        print(json.dumps(report, ensure_ascii=True), flush=True)
        return 0 if report["result"] in {"Draft", "Sent"} else 1
    except Exception as exc:
        report["result"] = "blocked_or_failed"
        report["error_type"] = type(exc).__name__
        # Do not dump API responses, tokens, or whole private mailbox objects.
        print(f"Live test blocked/failed: {type(exc).__name__}", flush=True)
        return 2
    finally:
        checkpoint()
        print(f"Test report: {output / 'result.json'}", flush=True)
        if driver and args.hold:
            input("Test browser is held for inspection. Press Enter to close this test browser only. ")
        if driver:
            engine.close_web_drivers()


if __name__ == "__main__":
    raise SystemExit(main())
