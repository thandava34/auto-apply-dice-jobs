import pandas as pd
import os
import datetime
import threading

COLUMNS = [
    "Date", "Posted Date", "Job Title", "Company", "Location", "Keywords", "Recruiter Name", "Recruiter Email",
    "Phone", "Status", "Resume Used", "Email Sent?",
    "Draft/Sent", "CC", "BCC", "Source", "Dedup Hash", "Job URL", "Description", "Phone Status"
]

class OutreachExcelStore:
    """
    Thread-safe Excel store with in-memory buffer.
    - Accumulates records in RAM; flushes to disk atomically.
    - avoids the O(N²) read-write-per-record anti-pattern.
    - Uses a lock so concurrent threads never corrupt the file.
    """

    def __init__(self, filepath="data/outreach_jobs.xlsx"):
        directory = os.path.dirname(filepath)
        if directory:
            os.makedirs(directory, exist_ok=True)
        self.filepath = filepath
        from utils.process_lock import ProcessLock
        self._lock = ProcessLock(filepath, timeout=10)

        # In-memory dedup set — populated from disk on first access
        self._known_hashes: set = set()
        self._hashes_loaded = False

        # Write buffer
        self._pending: list[dict] = []

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _ensure_file(self):
        """Create the xlsx if it doesn't exist yet."""
        if not os.path.exists(self.filepath):
            pd.DataFrame(columns=COLUMNS).to_excel(self.filepath, index=False)

    def _load_hashes(self):
        """Load all existing Dedup Hash values into memory (once)."""
        if self._hashes_loaded:
            return
        self._ensure_file()
        try:
            df = pd.read_excel(self.filepath, usecols=["Dedup Hash"], engine="openpyxl")
            self._known_hashes = set(df["Dedup Hash"].dropna().astype(str).tolist())
        except Exception:
            self._known_hashes = set()
        self._hashes_loaded = True

    def _save_dataframe_safely(self, df, filepath):
        """Atomic replacement; locked/corrupt destinations never become empty data."""
        import tempfile
        directory = os.path.dirname(os.path.abspath(filepath))
        fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=directory)
        os.close(fd)
        try:
            df.to_excel(temporary, index=False, engine='openpyxl')
            os.replace(temporary, filepath)
            return True
        finally:
            if os.path.exists(temporary):
                os.remove(temporary)

    def _flush_to_disk(self):
        if not self._pending:
            return
        # Any read failure propagates. Never reinterpret a failed read as empty.
        existing = pd.read_excel(self.filepath, engine='openpyxl') if os.path.exists(self.filepath) else pd.DataFrame(columns=COLUMNS)
        combined = pd.concat([existing, pd.DataFrame(self._pending)], ignore_index=True)
        for column in COLUMNS:
            if column not in combined:
                combined[column] = ''
        self._save_dataframe_safely(combined, self.filepath)
        self._pending.clear()

    def export_record(self, record):
        """Idempotent export retry; update one exact key or append a new row."""
        record = dict(record)
        if 'URL' in record:
            record.setdefault('Job URL', record.pop('URL'))
        with self._lock:
            self._flush_to_disk()
            df = pd.read_excel(self.filepath, engine='openpyxl') if os.path.exists(self.filepath) else pd.DataFrame(columns=COLUMNS)
            key = str(record.get('Dedup Hash') or '')
            if not key:
                raise ValueError('Export needs a stable record key')
            mask = df['Dedup Hash'].fillna('').astype(str).eq(key) if 'Dedup Hash' in df else pd.Series(False, index=df.index)
            if mask.sum() > 1:
                raise ValueError('Ambiguous duplicate export key; no workbook changes saved')
            if mask.any():
                for column, value in record.items():
                    if column == 'Phone Status':
                        continue  # An email export must not undo newer calling actions.
                    if column not in df:
                        df[column] = ''
                    df[column] = df[column].astype(object)
                    df.loc[mask, column] = value
            else:
                record.setdefault('Date', datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
                record.setdefault('Source', 'Nvoids')
                df = pd.concat([df, pd.DataFrame([record])], ignore_index=True)
            self._save_dataframe_safely(df, self.filepath)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def append_record(self, record: dict, auto_flush: bool = True):
        """
        Buffer a record and flush to disk efficiently.
        Keeps in-memory hash index updated instantly for zero-latency dedup checks.
        """
        with self._lock:
            self._load_hashes()
            record = dict(record)  # don't mutate caller's dict
            record["Date"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            # Normalise key names
            record.setdefault("Job URL", record.pop("URL", ""))
            record.setdefault("Source", "Nvoids")
            self._pending.append(record)
            
            # Update in-memory set so next is_in_excel() is instant (O(1))
            h = str(record.get("Dedup Hash", ""))
            if h:
                self._known_hashes.add(h)

            # Flush to disk every 10 items or when explicitly requested to eliminate O(N) disk I/O bottlenecks
            if auto_flush and len(self._pending) >= 10:
                self._flush_to_disk()

    def flush(self):
        """Force flush buffered records to disk instantly."""
        with self._lock:
            self._flush_to_disk()

    def is_in_excel(self, dedup_hash: str) -> bool:
        """O(1) duplicate check using in-memory set."""
        with self._lock:
            self._load_hashes()
            return str(dedup_hash) in self._known_hashes

    def has_listing(self, url):
        """Same Nvoids listing stays a duplicate even if its title is corrected."""
        from core.outreach.job_fields import listing_identity
        key = listing_identity(url)
        if not key:
            return False
        with self._lock:
            stamp = (os.path.getmtime(self.filepath), os.path.getsize(self.filepath)) if os.path.exists(self.filepath) else None
            if not hasattr(self, '_listing_stamp') or self._listing_stamp != stamp:
                keys = set()
                if stamp:
                    frame = pd.read_excel(self.filepath, engine='openpyxl').fillna('')
                    for column in ('Job URL', 'URL'):
                        if column in frame:
                            keys.update(listing_identity(v) for v in frame[column] if v)
                self._listing_keys, self._listing_stamp = keys, stamp
            return key in self._listing_keys or any(listing_identity(r.get('Job URL') or r.get('URL')) == key for r in self._pending)

    def update_record_status(self, dedup_hash: str, new_status: str, email_sent: str):
        """
        Update a single row's Status and Email Sent? fields.
        Reads once, patches, writes once.
        """
        with self._lock:
            if not os.path.exists(self.filepath):
                return
            try:
                df = pd.read_excel(self.filepath, engine="openpyxl")
                if "Dedup Hash" not in df.columns:
                    return
                mask = df["Dedup Hash"].astype(str) == str(dedup_hash)
                if mask.any():
                    df.loc[mask, "Status"] = new_status
                    df.loc[mask, "Email Sent?"] = email_sent
                    self._save_dataframe_safely(df, self.filepath)
            except Exception as e:
                print(f"[ExcelStore] update_record_status error: {e}")

    def update_phone_status(self, dedup_hash: str, status: str):
        with self._lock:
            self._flush_to_disk()
            df = pd.read_excel(self.filepath, engine='openpyxl')
            mask = df['Dedup Hash'].astype(str) == dedup_hash
            if mask.any():
                df['Phone Status'] = df.get('Phone Status', pd.Series('', index=df.index)).fillna('').astype(str)
                df.loc[mask, 'Phone Status'] = status
                self._save_dataframe_safely(df, self.filepath)

    def update_phone_records(self, records: list[dict], status: str):
        """All-or-nothing phone-only update; ambiguous historical keys never fan out."""
        from core.outreach.review import clean, needs_calling
        from openpyxl import load_workbook
        import tempfile
        if not records or not (status.startswith('Called (') or status == 'Skipped to Contact'):
            raise ValueError('Invalid phone action')
        with self._lock:
            self._flush_to_disk()
            workbook = load_workbook(self.filepath)
            temporary = None
            try:
                sheet = workbook.active
                headers = [cell.value for cell in sheet[1]]
                stored = [(index, dict(zip(headers, values))) for index, values in
                          enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)]
                targets = set()
                for record in records:
                    key = clean(record.get('Dedup Hash'))
                    if key:
                        matches = [(i, r) for i, r in stored if clean(r.get('Dedup Hash')) == key]
                    else:
                        phone, title = clean(record.get('Phone')), clean(record.get('Job Title'))
                        if not phone or not title:
                            raise ValueError('Record has no safe identifier; refresh and review it')
                        matches = [(i, r) for i, r in stored
                                   if clean(r.get('Phone')) == phone and clean(r.get('Job Title')) == title]
                    if len(matches) != 1 or not needs_calling(matches[0][1]):
                        raise ValueError('Record is ambiguous, missing or no longer eligible; no phone changes saved')
                    targets.add(matches[0][0])
                if 'Phone Status' not in headers:
                    headers.append('Phone Status')
                    sheet.cell(1, len(headers), 'Phone Status')
                column = headers.index('Phone Status') + 1
                for index in targets:
                    sheet.cell(index, column, status)
                fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=os.path.dirname(os.path.abspath(self.filepath)))
                os.close(fd)
                workbook.save(temporary)
                os.replace(temporary, self.filepath)
                temporary = None
                return len(targets)
            finally:
                workbook.close()
                if temporary and os.path.exists(temporary):
                    os.remove(temporary)

    def bulk_mark_email_status(self, target: str) -> list[dict]:
        """Mark all eligible email rows Draft or Sent and return changed records.

        This is a bookkeeping operation only. Rows without a usable recipient and
        rows already skipped/call-only are deliberately excluded.
        """
        normalized = str(target or "").strip().lower()
        if normalized not in {"draft", "sent"}:
            raise ValueError("target must be 'draft' or 'sent'")

        new_status = "Draft" if normalized == "draft" else "Sent"
        email_sent = "Drafted" if normalized == "draft" else "Yes"
        with self._lock:
            self._flush_to_disk()
            if not os.path.exists(self.filepath):
                return []

            df = pd.read_excel(self.filepath, engine="openpyxl")
            for column, default in (
                ("Recruiter Email", ""), ("Status", ""),
                ("Email Sent?", "No"), ("Draft/Sent", ""),
            ):
                if column not in df.columns:
                    df[column] = default

            recipients = df["Recruiter Email"].fillna("").astype(str).str.strip()
            statuses = df["Status"].fillna("").astype(str).str.strip()
            valid_recipient = recipients.str.match(r"^[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+$")
            excluded = statuses.str.contains(
                r"^(?:Skipped|No Email Found|Call Pending|Called|Skipped to Contact)",
                case=False,
                regex=True,
                na=False,
            )
            already_target = statuses.str.casefold().eq(new_status.casefold()) & (
                df["Email Sent?"].fillna("").astype(str).str.casefold().eq(email_sent.casefold())
            )
            mask = valid_recipient & ~excluded & ~already_target
            if not mask.any():
                return []

            changed = df.loc[mask].to_dict(orient="records")
            df.loc[mask, "Status"] = new_status
            df.loc[mask, "Email Sent?"] = email_sent
            df.loc[mask, "Draft/Sent"] = normalized
            self._save_dataframe_safely(df, self.filepath)
            return changed
