"""Read-only, bounded-page views of outreach records."""
from openpyxl import load_workbook
import re
from pathlib import Path


def clean(value):
    text = str(value or '').strip()
    return '' if text.casefold() in {'nan', 'none', 'nat'} else text


def email_outcome(record):
    status = clean(record.get('Status')).casefold()
    flag = clean(record.get('Email Sent?')).casefold()
    if status.startswith(('unconfirmed', 'review', 'needs review')) or flag == 'review':
        return 'Needs Review'
    if status.startswith(('error', 'failed')):
        return 'Failed'
    if status == 'sent':
        return 'Outcome unclear' if flag == 'drafted' else 'Sent'
    if status in {'draft', 'drafted'}:
        return 'Drafted'
    if flag == 'drafted':
        return 'Drafted'
    if flag in {'yes', 'sent'}:
        return 'Outcome unclear'
    if status.startswith(('pending', 'queued', 'skipped', 'call pending', 'no email')) or flag == 'no':
        return 'Not Drafted'
    return 'Outcome unclear'


def needs_calling(record):
    phone = clean(record.get('Phone'))
    handled = ('called', 'skipped')
    return (len(re.sub(r'\D', '', phone)) >= 7
            and not clean(record.get('Phone Status')).casefold().startswith(handled)
            and not clean(record.get('Status')).casefold().startswith(handled))


def records_page(path, page=0, size=50, search='', failed_only=False,
                 status_filter='All', calling_only=False, hide_called=False):
    search = clean(search)
    if page < 0 or (size is not None and not 1 <= size <= 100) or (size is None and page != 0):
        raise ValueError('Invalid page size or offset')
    if not Path(path).exists():
        return [], 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = workbook.active.iter_rows(values_only=True)
        headers = next(iterator, ())
        count, records = 0, []
        for values in iterator:
            record = dict(zip(headers, values))
            if not any(clean(value) for value in record.values()):
                continue
            if status_filter != 'All' and email_outcome(record) != status_filter:
                continue
            called = any(clean(record.get(key)).casefold().startswith('called')
                         for key in ('Phone Status', 'Status'))
            # The job browser hides completed calls until an explicit search.
            # Other callers (including bulk actions) keep their existing rules.
            search_history = hide_called and bool(search) and called
            if hide_called and called and not search:
                continue
            if calling_only and not needs_calling(record) and not search_history:
                continue
            if failed_only and not str(record.get('Status', '')).startswith(('Error:', 'Unconfirmed:')):
                continue
            if search and search.casefold() not in ' '.join(str(record.get(k) or '') for k in (
                    'Recruiter Email', 'Company', 'Job Title', 'Phone', 'Recruiter Name')).casefold():
                continue
            if size is None or page * size <= count < (page + 1) * size:
                records.append(record)
            count += 1
        return records, count
    finally:
        workbook.close()
